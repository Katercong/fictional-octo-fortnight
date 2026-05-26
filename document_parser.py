"""
文档解析模块
支持 TXT、PDF、DOCX 格式的文档解析
"""

import os
from datetime import datetime
from PyPDF2 import PdfReader
from docx import Document

def extract_text_from_txt(file_path: str) -> str:
    """
    提取 TXT 文本文件内容
    
    Args:
        file_path: 文件路径
        
    Returns:
        文本内容，如果失败则返回错误信息
    """
    try:
        # 使用 UTF-8 编码读取，遇到错误时替换
        with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
            return f.read().strip()
    except Exception as e:
        return f"读取TXT失败: {str(e)}"

def extract_text_from_pdf(file_path: str) -> str:
    """
    提取 PDF 文件内容
    
    Args:
        file_path: 文件路径
        
    Returns:
        文本内容，如果失败则返回错误信息
    """
    try:
        with open(file_path, 'rb') as f:
            reader = PdfReader(f)
            text = ""
            # 逐页提取文本
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
            return text.strip()
    except Exception as e:
        return f"读取PDF失败: {str(e)}"

def extract_text_from_docx(file_path: str) -> str:
    """
    提取 DOCX (Word) 文件内容
    
    Args:
        file_path: 文件路径
        
    Returns:
        文本内容，如果失败则返回错误信息
    """
    try:
        doc = Document(file_path)
        # 提取所有段落文本，用换行符连接
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        return text.strip()
    except Exception as e:
        return f"读取Word失败: {str(e)}"

def extract_text_from_file(file_path: str) -> str:
    """
    根据文件扩展名自动选择解析器提取文本
    
    Args:
        file_path: 文件路径
        
    Returns:
        文本内容
    """
    filename = file_path.lower()
    
    if filename.endswith('.txt'):
        return extract_text_from_txt(file_path)
    elif filename.endswith('.pdf'):
        return extract_text_from_pdf(file_path)
    elif filename.endswith('.docx'):
        return extract_text_from_docx(file_path)
    else:
        return f"不支持的文件格式"

def extract_text_from_uploadfile_pdf(file) -> str:
    """
    从 FastAPI UploadFile 提取 PDF 内容（用于上传接口）
    
    Args:
        file: FastAPI UploadFile 对象
        
    Returns:
        文本内容
    """
    try:
        reader = PdfReader(file.file)
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text.strip()
    except Exception as e:
        return f"读取PDF失败: {str(e)}"

def extract_text_from_uploadfile_docx(file) -> str:
    """
    从 FastAPI UploadFile 提取 DOCX 内容（用于上传接口）
    
    Args:
        file: FastAPI UploadFile 对象
        
    Returns:
        文本内容
    """
    try:
        doc = Document(file.file)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        return text.strip()
    except Exception as e:
        return f"读取Word失败: {str(e)}"

def extract_text_from_uploadfile_txt(file) -> str:
    """
    从 FastAPI UploadFile 提取 TXT 内容（用于上传接口）
    
    Args:
        file: FastAPI UploadFile 对象
        
    Returns:
        文本内容
    """
    try:
        content = file.file.read().decode('utf-8', errors='replace')
        return content.strip()
    except Exception as e:
        return f"读取TXT失败: {str(e)}"

def extract_text_from_uploadfile(file) -> str:
    """
    从 FastAPI UploadFile 自动选择解析器提取文本（用于上传接口）
    
    Args:
        file: FastAPI UploadFile 对象
        
    Returns:
        文本内容
    """
    filename = file.filename.lower()

    if filename.endswith('.pdf'):
        return extract_text_from_uploadfile_pdf(file)
    elif filename.endswith('.docx'):
        return extract_text_from_uploadfile_docx(file)
    elif filename.endswith('.txt'):
        return extract_text_from_uploadfile_txt(file)
    else:
        return f"不支持的文件格式: {filename}"

def parse_documents_in_folder(folder_path: str) -> list:
    """
    解析文件夹中的所有支持的文档
    
    Args:
        folder_path: 文件夹路径
        
    Returns:
        文档列表
    """
    result = []
    supported_extensions = ('.txt', '.pdf', '.docx')
    
    # 验证路径
    if not os.path.isdir(folder_path):
        return [{"error": f"路径不是有效的文件夹: {folder_path}"}]
    
    # 递归遍历
    for root, _, files in os.walk(folder_path):
        for filename in files:
            if filename.lower().endswith(supported_extensions):
                file_path = os.path.join(root, filename)
                
                try:
                    # 获取文件修改时间
                    file_stat = os.stat(file_path)
                    modify_time = datetime.fromtimestamp(file_stat.st_mtime).strftime('%Y-%m-%d %H:%M:%S')
                    
                    # 提取内容
                    content = extract_text_from_file(file_path)
                    
                    result.append({
                        "filename": filename,
                        "file_path": file_path,
                        "content": content,
                        "modify_time": modify_time
                    })
                except Exception as e:
                    result.append({
                        "filename": filename,
                        "file_path": file_path,
                        "content": f"处理文件时发生错误: {str(e)}",
                        "modify_time": "未知"
                    })
    
    return result

if __name__ == "__main__":
    """
    命令行测试：解析指定文件夹并显示结果
    """
    import sys
    
    if len(sys.argv) != 2:
        print("用法: python document_parser.py <文件夹路径>")
        sys.exit(1)
    
    folder_path = sys.argv[1]
    documents = parse_documents_in_folder(folder_path)
    
    print(f"共找到 {len(documents)} 个文档:")
    for doc in documents:
        print(f"\n文件名: {doc['filename']}")
        print(f"路径: {doc['file_path']}")
        print(f"修改时间: {doc['modify_time']}")
        print(f"内容预览: {doc['content'][:200]}...")
